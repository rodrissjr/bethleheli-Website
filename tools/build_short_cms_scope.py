from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("deliverables/Betheli_Sanity_CMS_Field_Scope_Short.docx").resolve()
INK = "243138"
MUTED = "66737C"
NAVY = "243B49"
PALE = "F2F5F6"
ACCENT = "EA4729"
GRID = "D9D9D9"


def font(run, size=10.5, bold=False, color=INK, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font(r, size=17 if level == 1 else 12.5, bold=True, color="000000", name="Aptos Display")
    return p


def lead(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label + " ")
    font(r, bold=True, color=INK)
    r = p.add_run(text)
    font(r, color=INK)


def compact_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.42)
    table.columns[1].width = Inches(5.45)
    headers = table.rows[0].cells
    for index, text in enumerate(("Content type", "Required fields")):
        headers[index].width = table.columns[index].width
        headers[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_fill(headers[index], NAVY)
        cell_border(headers[index])
        cell_margins(headers[index])
        p = headers[index].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        font(r, size=9.5, bold=True, color="FFFFFF")

    for row_index, (name, fields) in enumerate(rows):
        cells = table.add_row().cells
        for index, text in enumerate((name, fields)):
            cells[index].width = table.columns[index].width
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_fill(cells[index], PALE if row_index % 2 else "FFFFFF")
            cell_border(cells[index])
            cell_margins(cells[index])
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            r = p.add_run(text)
            font(r, size=9.2, bold=index == 0, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def setup(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.66)
    section.right_margin = Inches(0.66)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style_name == "Title":
            p_pr = style.element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("BETHELI EXPEDITIONS  SANITY CMS")
    font(r, size=8, bold=True, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("CMS field scope  |  ")
    font(r, size=8, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_end)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup(doc)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    title_pr = title._p.get_or_add_pPr()
    border = title_pr.find(qn("w:pBdr"))
    if border is not None:
        title_pr.remove(border)
    r = title.add_run("Betheli Sanity CMS Field Scope")
    font(r, size=25, bold=True, color="000000", name="Aptos Display")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    r = subtitle.add_run("Short handoff for Trips Itineraries Accommodations and Group Climbs")
    font(r, size=11.5, bold=True, color=ACCENT)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.line_spacing = 1.1
    r = intro.add_run(
        "This is the approved CMS scope. Every listed field is required. The Tanzania map is stored once and reused by every trip. Day numbers, filters, sliders and calculated labels remain in the frontend."
    )
    font(r, size=10.5, color=INK)

    heading(doc, "Trips and Itineraries", 1)
    compact_table(doc, [
        ("Trip Settings", "map: image asset and alternative text"),
        ("Trip Identity", "title, card title, slug, category, location label, duration in days, starting price, card image"),
        ("Trip Hero", "hero image, hero summary"),
        ("Trip Overview", "overview paragraphs, overview image, itinerary-for text, travel-dates text, guests text"),
        ("Trip Planning", "destinations, planning title, planning description"),
        ("Trip Content", "ordered itinerary days, landscapes, inclusion groups, related trip references"),
        ("Trip SEO", "SEO title, description, canonical path, social sharing image"),
        ("Itinerary Day", "title, summary, experience image, experience description, meal plan, accommodation references"),
    ])

    heading(doc, "Accommodation", 1)
    compact_table(doc, [
        ("Accommodation", "property name, slug, tier, style, destination group, location, description, cover image, gallery images"),
        ("Tier values", "Budget, Mid-range, Luxury"),
        ("Relationship", "Each itinerary day stores an ordered list of Accommodation references"),
    ])
    lead(doc, "Image rule.", "Every image requires alternative text. Accommodation gallery order controls the carousel order.")

    doc.add_page_break()
    heading(doc, "Group Climbs", 1)
    compact_table(doc, [
        ("Group Climbs Page", "hero eyebrow, title, description, hero image, benefits, schedule eyebrow, schedule title, empty-state title, empty-state description, SEO"),
        ("Benefit", "icon key, title, description"),
        ("Group Departure", "route name, duration in days, start date, end date, highlights, price, deposit, booking status, promotion status, season status"),
        ("Discounted price", "original price is required when promotion status is Discounted Price"),
        ("Rain season", "season label is required when season status is Rain Season"),
    ])

    heading(doc, "Shared Objects", 1)
    compact_table(doc, [
        ("Content Image", "image asset, alternative text, hotspot and crop"),
        ("Money", "amount and USD currency"),
        ("SEO", "title, description, canonical path and social sharing image"),
        ("Landscape", "name, descriptor, icon key, image and description"),
        ("Inclusion Group", "marker, title, description and inclusion items"),
        ("Inclusion Item", "icon key, title and description"),
    ])

    heading(doc, "Generated by the Website", 1)
    compact_table(doc, [
        ("Trip interface", "day numbers, accordion state, map modal, slider position, image count and 4.5-second timing"),
        ("Trip listing", "category filters, duration filters, visible trip count and category counts"),
        ("Group schedule", "year and month filters, date-range text, route title, discount percentage and join-form selection"),
    ])
    lead(doc, "Validation rule.", "Publishing is blocked when a required field is missing. Trip itinerary length must equal the trip duration, and a departure end date must be later than its start date.")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
