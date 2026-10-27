from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/home/iamrodriss/Desktop/code/betleheli")
OUTPUT = ROOT / "deliverables" / "Betheli_Sanity_CMS_Required_Fields.docx"

INK = "26363D"
NAVY = "243844"
ACCENT = "F04B32"
PALE = "F3F6F7"
PALE_BLUE = "EAF0F3"
BORDER = "D9D9D9"
MUTED = "617078"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p_pr = p._p.get_or_add_pPr()
    p_borders = p_pr.find(qn("w:pBdr"))
    if p_borders is not None:
        p_pr.remove(p_borders)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(54)
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run("Betheli Expeditions Sanity CMS Required Field Specification")
    set_run_font(run, size=28, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(30)
    r = subtitle.add_run("Existing website content only")
    set_run_font(r, size=15, bold=True, color=ACCENT)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(20)
    intro.paragraph_format.line_spacing = 1.18
    r = intro.add_run(
        "This document defines the exact Sanity fields required to replace the content that is currently hardcoded in the Betheli Expeditions website. "
        "It covers the Trips listing, trip itinerary page, accommodations, Kilimanjaro landing and route pages, Group Climbs, group departures, FAQs, reviews, the local advisor, global contact details and page SEO."
    )
    set_run_font(r, size=12, color=INK)

    decision = doc.add_paragraph()
    decision.paragraph_format.space_after = Pt(30)
    decision.paragraph_format.line_spacing = 1.18
    r = decision.add_run(
        "Every field listed in this specification is mandatory for its document type. The schema must prevent publishing when a required value is missing. "
        "Values identified as generated are calculated by the website and must not be entered twice in Sanity."
    )
    set_run_font(r, size=12, bold=True, color=INK)

    meta = [
        ("Sanity project", "ife51h63"),
        ("Dataset", "production"),
        ("Frontend", "Next.js App Router"),
        ("Studio", "Standalone Studio in studio"),
        ("Prepared", "17 October 2026"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = [Inches(1.65), Inches(4.75)]
    for i, (label, value) in enumerate(meta):
        for j, text in enumerate((label, value)):
            cell = table.cell(i, j)
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 90, 110, 90, 110)
            set_cell_borders(cell)
            if j == 0:
                shade(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, bold=(j == 0), color=INK)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(7)
    for run in p.runs:
        set_run_font(run, size=19 if level == 1 else 13, bold=True, color="000000")


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.space_after = Pt(7)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, size=10.5, bold=True, color=INK)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r, size=10.5, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.3, color=INK)


def add_field_table(doc: Document, rows: list[tuple[str, str, str]], widths=(1.7, 1.55, 3.95)) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ("Sanity field", "Sanity type", "Exact use on the current site")
    header = table.rows[0]
    set_repeat_table_header(header)
    for j, text in enumerate(headers):
        cell = header.cells[j]
        cell.width = Inches(widths[j])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        set_cell_margins(cell, 115, 125, 115, 125)
        set_cell_borders(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.3, bold=True, color="FFFFFF")

    for i, row_data in enumerate(rows):
        row = table.add_row()
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell, "FFFFFF" if i % 2 == 0 else PALE)
            set_cell_margins(cell, 95, 120, 95, 120)
            set_cell_borders(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, size=9.2, bold=(j == 0), color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def set_document_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style_name == "Title":
            p_pr = style.element.get_or_add_pPr()
            p_borders = p_pr.find(qn("w:pBdr"))
            if p_borders is not None:
                p_pr.remove(p_borders)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.72)
        section.bottom_margin = Inches(0.68)
        section.left_margin = Inches(0.66)
        section.right_margin = Inches(0.66)
        section.header_distance = Inches(0.3)
        section.footer_distance = Inches(0.3)


def add_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("BETHELI EXPEDITIONS   SANITY CMS")
    set_run_font(r, size=8, bold=True, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Required field specification")
    set_run_font(run, size=8, color=MUTED)
    run = p.add_run("   |   ")
    set_run_font(run, size=8, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def build() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_document_styles(doc)
    configure_sections(doc)
    add_header_footer(doc)
    add_title_page(doc)

    add_page_break(doc)
    add_heading(doc, "Content Model Decision", 1)
    add_body(doc, "The website needs fourteen Sanity document types and four shared object types. This is the smallest model that covers the content currently displayed without storing presentation details such as card width, slider speed or CSS classes.")
    add_heading(doc, "Sanity document types", 2)
    add_bullets(doc, [
        "Site Settings singleton",
        "Trips Page singleton",
        "Trip",
        "Accommodation",
        "Accommodation Collection",
        "Kilimanjaro Page singleton",
        "Kilimanjaro Shared Content singleton",
        "Kilimanjaro Route",
        "Group Climbs Page singleton",
        "Group Departure",
        "FAQ Category",
        "FAQ",
        "Testimonial",
        "Local Advisor",
    ])
    add_heading(doc, "Shared object types", 2)
    add_field_table(doc, [
        ("contentImage", "object", "image asset, alternative text and hotspot or crop data"),
        ("seo", "object", "SEO title, meta description, canonical path and social sharing image"),
        ("callToAction", "object", "button label and destination path used by the visible page CTA"),
        ("price", "object", "numeric amount, currency code and displayed unit such as per person"),
    ])
    add_body(doc, "Array order is the display order. Sanity generates an internal key for every array item; the frontend must query and use that key instead of using the item index.")

    add_page_break(doc)
    add_heading(doc, "Rules Applied to Every Public Page", 1)
    add_field_table(doc, [
        ("title", "string", "Public page or content title displayed on the website"),
        ("slug", "slug", "Lowercase URL segment with hyphens; unique within its document type"),
        ("seo.title", "string", "Browser and search title returned by generateMetadata"),
        ("seo.description", "text", "Meta description and social description"),
        ("seo.canonicalPath", "string", "Canonical path such as /trips/northern-serengeti"),
        ("seo.shareImage", "contentImage", "Open Graph and Twitter image with alternative text"),
    ])
    add_body(doc, "All image fields use Sanity image assets with hotspot enabled. Every image requires alternative text. Decorative interface icons remain in code and are not CMS fields.")

    add_heading(doc, "Site Settings", 1)
    add_body(doc, "Document ID: siteSettings. One document controls the global metadata, contact details, review summary and footer information already shown across the site.")
    add_field_table(doc, [
        ("siteName", "string", "Betheli Expeditions"),
        ("siteUrl", "url", "Production base URL used by metadata, canonical URLs and the sitemap"),
        ("defaultTitle", "string", "Default browser title currently used in the root layout"),
        ("titleTemplate", "string", "Page title template ending with Betheli Expeditions"),
        ("defaultDescription", "text", "Default site description used when rendering global metadata"),
        ("logo", "contentImage", "Header and footer company logo"),
        ("defaultShareImage", "contentImage", "Fallback social image"),
        ("email", "string", "info@betheliexpeditions.com"),
        ("phone", "string", "+255 754 411 965"),
        ("whatsappNumber", "string", "+255 692 796 412"),
        ("officeAddress", "string", "Saba Saba, Arusha, Tanzania"),
        ("addressLocality", "string", "Arusha for TourOperator structured data"),
        ("addressCountry", "string", "TZ for TourOperator structured data"),
        ("areaServed", "array of string", "Tanzania and Zanzibar"),
        ("footerDescription", "text", "Tanzanian-owned operator description in the footer"),
        ("copyrightText", "string", "Copyright and operator registration statement"),
        ("reviewRating", "number", "5.0 review summary displayed in review and hero sections"),
        ("reviewCount", "number", "23 verified reviews displayed on the site"),
        ("tripadvisorUrl", "url", "Company review profile destination"),
    ])

    add_heading(doc, "Footer Navigation", 2)
    add_field_table(doc, [
        ("footerColumns", "array of object", "CLIMBS, SAFARIS and EXPLORE columns in visible order"),
        ("footerColumns[].heading", "string", "Column heading"),
        ("footerColumns[].links", "array of callToAction", "Visible link labels and internal destinations"),
        ("credentials", "array of object", "TATO Member, ATTA Member, AMREF Covered and 5.0 Rated rows"),
        ("credentials[].iconKey", "string", "Frontend icon identifier"),
        ("credentials[].title", "string", "Credential name"),
        ("credentials[].description", "string", "Credential explanation"),
    ])

    add_heading(doc, "Trips Page", 1)
    add_body(doc, "Document ID: tripsPage. Route: /trips. The trip cards and filter values are generated from Trip documents; they are not entered again on this page.")
    add_heading(doc, "Hero", 2)
    add_field_table(doc, [
        ("hero.eyebrow", "string", "TAILORED TANZANIA JOURNEYS"),
        ("hero.title", "string", "Trips across Tanzania"),
        ("hero.lead", "string", "Climb, safari and coast lead sentence"),
        ("hero.description", "text", "Browse a selection of trip ideas description"),
        ("hero.image", "contentImage", "Trips hero background"),
        ("hero.primaryAction", "callToAction", "Start planning button"),
        ("hero.secondaryAction", "callToAction", "Explore trips anchor"),
    ])
    add_heading(doc, "Listing and Closing CTA", 2)
    add_field_table(doc, [
        ("listing.eyebrow", "string", "FIND YOUR TANZANIA JOURNEY"),
        ("listing.title", "string", "Kilimanjaro, Safari and Zanzibar"),
        ("listing.description", "text", "Choose a starting point description"),
        ("listing.emptyMessage", "string", "No trips match these filters message"),
        ("closingCta.title", "string", "Not sure which trip fits"),
        ("closingCta.description", "text", "Prompt asking for places, time and travellers"),
        ("closingCta.image", "contentImage", "Closing CTA background image"),
        ("seo", "seo", "Metadata for /trips"),
    ])
    add_body(doc, "Generated by the frontend: All, Zanzibar, Safari and Kilimanjaro filters; duration options; visible trip count; category counts; clear-filter state; trip cards.", bold_lead="Generated by the frontend:")

    add_page_break(doc)
    add_heading(doc, "Trip Document and Itinerary Page", 1)
    add_body(doc, "Document type: trip. Current route example: /trips/northern-serengeti. Each Trip document supplies its own listing card, hero, overview, itinerary, accommodations, landscapes, inclusions, related trips, reviews, FAQs and SEO.")
    add_heading(doc, "Identity and Listing Card", 2)
    add_field_table(doc, [
        ("title", "string", "6-Day Northern Serengeti River Crossing Safari"),
        ("cardTitle", "string", "Northern Serengeti River Crossing"),
        ("slug", "slug", "northern-serengeti"),
        ("category", "string list", "Safari, Zanzibar or Kilimanjaro"),
        ("locationLabel", "string", "Location value used by homepage filters"),
        ("durationDays", "number", "6, 8 or 10 as displayed on trip cards"),
        ("startingPrice", "price", "Amount, USD currency and per person unit"),
        ("cardImage", "contentImage", "Image and alt text used on listing cards"),
    ])
    add_heading(doc, "Hero and Overview", 2)
    add_field_table(doc, [
        ("heroImage", "contentImage", "Full-width trip hero image"),
        ("heroSummary", "text", "Migration-focused safari summary under the page title"),
        ("overviewTitle", "string", "Trip title repeated in the overview section"),
        ("overviewParagraphs", "array of text", "The two visible overview paragraphs"),
        ("overviewImage", "contentImage", "Hot-air balloons over the Serengeti image"),
        ("facts", "array of object", "Itinerary for, Travel Dates and Guests rows"),
        ("facts[].label", "string", "Visible fact label"),
        ("facts[].value", "string", "Visible fact value"),
        ("facts[].iconKey", "string", "route, calendar or guests icon identifier"),
    ])
    add_heading(doc, "Map and Planning Panel", 2)
    add_field_table(doc, [
        ("mapImage", "contentImage", "Tanzania national parks map and accessible description"),
        ("destinations", "array of string", "Tarangire, Ngorongoro, Central Serengeti and Northern Serengeti Mara River"),
        ("planningPanel.label", "string", "Plan Trip"),
        ("planningPanel.title", "string", "Experience the 6-Day Northern Serengeti Safari for yourself"),
        ("planningPanel.description", "text", "Senior planner proposal message"),
        ("planningPanel.quoteAction", "callToAction", "Request a Quote"),
        ("planningPanel.whatsappLabel", "string", "WhatsApp Us"),
    ])

    add_heading(doc, "Itinerary Days", 2)
    add_field_table(doc, [
        ("itinerary", "array of itineraryDay", "Ordered six-day itinerary"),
        ("itinerary[].title", "string", "Arrival, Tarangire National Park, Ngorongoro Crater and the remaining day titles"),
        ("itinerary[].summary", "text", "Paragraph displayed immediately after the day heading"),
        ("itinerary[].image", "contentImage", "Day experience image with alternative text"),
        ("itinerary[].experience", "text", "Longer experience paragraph displayed below the image"),
        ("itinerary[].mealPlan", "string list", "Half Board, Full Board or Breakfast"),
        ("itinerary[].accommodations", "array of reference", "References to the visible Mid-range and Luxury choices"),
    ])
    add_body(doc, "Generated by the frontend: Day 1, Day 2 and subsequent numbering; Show All Days state; slider image count; automatic 4.5-second image change; previous and next controls.", bold_lead="Generated by the frontend:")

    add_heading(doc, "Landscapes Inclusions and Related Content", 2)
    add_field_table(doc, [
        ("landscapes", "array of object", "Tarangire, Ngorongoro, Central Serengeti, Northern Serengeti and Mara River tabs"),
        ("landscapes[].icon", "string", "Visible emoji or icon identifier"),
        ("landscapes[].title", "string", "Landscape name"),
        ("landscapes[].descriptor", "string", "Baobabs and elephants, Crater floor and similar label"),
        ("landscapes[].image", "contentImage", "Landscape tab image and alt text"),
        ("landscapes[].description", "text", "Expanded landscape description"),
        ("inclusionIntro.eyebrow", "string", "YOUR SAFARI, TAKEN CARE OF"),
        ("inclusionIntro.title", "string", "What is included in your safari"),
        ("inclusionIntro.description", "text", "Quote-confirmation message"),
        ("inclusionGroups", "array of object", "Preparation, Safari Days and Ongoing Care groups"),
        ("inclusionGroups[].marker", "string", "01 / PREPARATION and subsequent markers"),
        ("inclusionGroups[].title", "string", "Group heading"),
        ("inclusionGroups[].description", "text", "Group supporting sentence"),
        ("inclusionGroups[].items", "array of object", "Visible included-service items"),
        ("inclusionGroups[].items[].iconKey", "string", "Frontend icon identifier"),
        ("inclusionGroups[].items[].title", "string", "Included-service title"),
        ("inclusionGroups[].items[].description", "text", "Included-service explanation"),
        ("testimonials", "array of reference", "Reviews displayed in the trip review slider"),
        ("relatedTrips", "array of reference", "You Might Also Like trip cards"),
        ("faqs", "array of reference", "Questions displayed in the trip FAQ section"),
        ("closingCta", "object", "Title, description and images used by the visible CTA banner"),
        ("seo", "seo", "Metadata for the trip route"),
    ])

    add_heading(doc, "Accommodation Document", 1)
    add_body(doc, "Document type: accommodation. One document represents one real property. It is referenced by itinerary days and by accommodation collections, so changing an image or tier updates every place where that property appears.")
    add_field_table(doc, [
        ("propertyName", "string", "Karibu River Camp, Legendary Serengeti, Mount Meru Hotel and every other named property"),
        ("slug", "slug", "Stable property identifier used by Studio references"),
        ("tier", "string list", "Budget, Mid-range or Luxury badge shown over the image"),
        ("destinationGroup", "string list", "Safari Mainland or Zanzibar, matching the existing property libraries"),
        ("summary", "text", "Property description shown when that property is the featured accommodation"),
        ("coverImage", "contentImage", "Main itinerary or selector image with alt text"),
        ("gallery", "array of contentImage", "Ordered property images used by the accommodation carousel"),
    ])
    add_body(doc, "Classification rule: tier is selected by the content editor from Budget, Mid-range or Luxury. The website must not infer Budget merely because a property is outside a national park.")

    add_heading(doc, "Accommodation Collection", 1)
    add_body(doc, "Document type: accommodationCollection. This represents the visible choices such as Classic safari camp, Tented suite, Forest lodge, Arusha garden lodge, Resort stay, Beachfront villa, Private suite and Family-friendly resort.")
    add_field_table(doc, [
        ("title", "string", "Visible collection name such as Classic safari camp"),
        ("destinationGroup", "string list", "Safari Mainland or Zanzibar"),
        ("description", "text", "Visible explanation beneath the collection title"),
        ("featuredAccommodation", "reference", "Property shown first when the collection is selected"),
        ("accommodations", "array of reference", "Ordered property images available in the slider"),
        ("displayOrder", "number", "Order of the four visible collection choices"),
    ])
    add_body(doc, "The array length is not forced to ten. The slider displays every published referenced property available for that collection and tier.")

    add_heading(doc, "Kilimanjaro Page", 1)
    add_body(doc, "Document ID: kilimanjaroPage. Route: /kilimanjaro. Route cards and the comparison table are generated from Kilimanjaro Route documents so prices, days and difficulty are entered once.")
    add_field_table(doc, [
        ("hero.eyebrow", "string", "TANZANIA THE ROOF OF AFRICA"),
        ("hero.title", "string", "Climb Mount Kilimanjaro"),
        ("hero.description", "text", "Experienced local guides and carefully planned routes message"),
        ("hero.image", "contentImage", "Kilimanjaro hero background"),
        ("hero.primaryAction", "callToAction", "Plan My Climb"),
        ("hero.secondaryAction", "callToAction", "Explore Routes"),
        ("heroTrustItems", "array of string", "Local guides, safety, personal support and small groups"),
        ("introduction.eyebrow", "string", "THE MOUNTAIN, AND EVERYTHING BETWEEN"),
        ("introduction.title", "string", "More Than a Summit"),
        ("introduction.description", "text", "Journey through the mountain environments"),
        ("introduction.body", "array of text", "The two visible supporting statements"),
        ("introduction.image", "contentImage", "Climbers above the clouds image"),
        ("facts", "array of object", "5,895 m, ecological zones, routes and unforgettable experience"),
        ("facts[].iconKey", "string", "Frontend icon identifier"),
        ("facts[].value", "string", "Main fact value"),
        ("facts[].description", "string", "Fact explanation"),
        ("whySection.title", "string", "Why Climb Kilimanjaro With Us"),
        ("whySection.image", "contentImage", "Section image"),
        ("whySection.items", "array of object", "Six current reasons with title and description"),
        ("routesIntro", "object", "Eyebrow, title and description above route cards"),
        ("journeyIntro", "object", "Eyebrow, title and description above the five journey stages"),
        ("journeyStages", "array of object", "Rainforest, Moorland, Alpine Desert, Summit and Descent"),
        ("journeyStages[].stageLabel", "string", "Number and environment label"),
        ("journeyStages[].title", "string", "Stage headline"),
        ("journeyStages[].description", "text", "Stage description"),
        ("journeyStages[].image", "contentImage", "Stage image and alt text"),
        ("travelStyles", "array of object", "Solo, Private Group and Join a Group cards"),
        ("climbExperienceSteps", "array of object", "Training through Descent and Celebration"),
        ("groupDepartureRoutes", "array of reference", "Four Join a group cards"),
        ("testimonials", "array of reference", "Reviews displayed on the landing page"),
        ("faqs", "array of reference", "Kilimanjaro FAQ preview"),
        ("closingCta", "object", "Ready to Climb Kilimanjaro title and image"),
        ("seo", "seo", "Metadata for /kilimanjaro"),
    ])

    add_heading(doc, "Kilimanjaro Route Document", 1)
    add_body(doc, "Document type: kilimanjaroRoute. It drives Machame, Lemosho, Marangu, Rongai, Northern Circuit and Umbwe pages and also supplies the route cards and comparison table on /kilimanjaro.")
    add_heading(doc, "Route Identity and Hero", 2)
    add_field_table(doc, [
        ("name", "string", "Full route name such as Lemosho Route"),
        ("shortName", "string", "Short name used in CTA text"),
        ("slug", "slug", "Route URL segment"),
        ("durationLabel", "string", "7-8 Days or the route-specific visible range"),
        ("itineraryDays", "number", "Number of itinerary records rendered"),
        ("difficulty", "string list", "Moderate-hard, Hard, Challenging or Very challenging as used by the site"),
        ("scenery", "string", "Exceptional, Classic, Remote, Panoramic or Dramatic"),
        ("crowdLevel", "string", "Value displayed in the landing-page comparison table"),
        ("startingPrice", "price", "Numeric amount used on route cards and the comparison table"),
        ("heroLead", "string", "Short sentence directly below the route name"),
        ("heroDescription", "text", "Route-specific hero paragraph"),
        ("heroImage", "contentImage", "Hero and route landscape image"),
        ("mapImage", "contentImage", "Illustrated map and alt text"),
    ])
    add_heading(doc, "Overview Benefits and Route Map", 2)
    add_field_table(doc, [
        ("overviewParagraphs", "array of text", "The two visible route overview paragraphs"),
        ("bestFor", "string", "Best for value in the route facts"),
        ("startingPoint", "string", "Trailhead or gate displayed in route facts"),
        ("highestPoint", "string", "5,895 m value displayed in route facts"),
        ("accommodationStyle", "string", "Camping route or shared mountain huts"),
        ("benefits", "array of object", "Four numbered Why Choose this Route cards"),
        ("benefits[].title", "string", "Benefit heading"),
        ("benefits[].description", "text", "Benefit explanation"),
        ("mapSteps", "array of string", "Four step-by-step route map statements"),
        ("mapNote", "text", "Illustrative route only message"),
    ])
    add_heading(doc, "Route Itinerary", 2)
    add_field_table(doc, [
        ("itineraryIntro.eyebrow", "string", "ONE MOUNTAIN and written number of days"),
        ("itineraryIntro.title", "string", "Route name plus Itinerary"),
        ("itineraryIntro.description", "text", "Day-by-day outline description"),
        ("itineraryNote", "text", "Route-specific scheduling note"),
        ("itinerary", "array of routeDay", "Ordered route days"),
        ("itinerary[].title", "string", "Day heading"),
        ("itinerary[].location", "string", "From-to route label"),
        ("itinerary[].description", "text", "Day description"),
        ("itinerary[].distance", "string", "Visible distance text; use To confirm until verified"),
        ("itinerary[].walkingTime", "string", "Visible walking-time text; use To confirm until verified"),
        ("itinerary[].elevation", "string", "Visible elevation text; use To confirm until verified"),
        ("itinerary[].overnight", "string", "Visible overnight text; use To confirm until verified"),
        ("itinerary[].image", "contentImage", "Day image and alt text"),
    ])
    add_body(doc, "Generated by the frontend: Day numbering, route comparison row, itinerary array count and related-route exclusion.", bold_lead="Generated by the frontend:")

    add_heading(doc, "Route Supporting Sections", 2)
    add_field_table(doc, [
        ("inclusionGroups", "array of object", "Preparation, On the Mountain and Ongoing Care groups with marker, title, description and items"),
        ("summitNight.title", "string", "Summit Night"),
        ("summitNight.heading", "string", "Slow steps, cold air and sunrise statement"),
        ("summitNight.paragraphs", "array of text", "The two visible summit-night paragraphs"),
        ("packingIntro", "object", "Eyebrow, title and description for What Should You Bring"),
        ("packingItems", "array of object", "The six visible packing groups, from Clothing through the final Gear group"),
        ("safetyContent", "reference", "Shared safety block displayed on the route"),
        ("fitIntro", "object", "Is This Route Right For You heading and route-specific description"),
        ("goodFor", "array of string", "Four route-specific traveller types"),
        ("considerCarefully", "array of object", "Altitude, Long summit day, Weather and Physical demands"),
        ("testimonials", "array of reference", "Route review slider"),
        ("relatedRoutes", "array of reference", "Compare other routes cards"),
        ("faqs", "array of reference", "Ten route-specific questions"),
        ("closingCta", "object", "Ready to climb route title and image"),
        ("seo", "seo", "Metadata for the individual route"),
    ])

    add_heading(doc, "Kilimanjaro Shared Content", 2)
    add_body(doc, "Document ID: kilimanjaroSharedContent. The safety section is the same component across the Kilimanjaro landing and route pages, so the content is stored once and referenced everywhere.")
    add_field_table(doc, [
        ("safety.eyebrow", "string", "CARE BEFORE THE SUMMIT"),
        ("safety.title", "string", "Your Safety Comes First"),
        ("safety.description", "text", "Mountain safety introduction"),
        ("safety.image", "contentImage", "Experienced mountain guides image"),
        ("safety.statement", "text", "Returning safely is more important than the summit statement"),
        ("safety.items", "array of object", "Each visible safety topic with icon key, title and explanation"),
        ("safety.action", "callToAction", "Talk to a Guide"),
    ])

    add_heading(doc, "Group Climbs Page", 1)
    add_body(doc, "Document ID: groupClimbsPage. Route: /groups. Years and months are generated from Group Departure dates rather than maintained as page fields.")
    add_field_table(doc, [
        ("hero.eyebrow", "string", "CLIMB TOGETHER"),
        ("hero.title", "string", "Share the journey. Reach the Summit together"),
        ("hero.description", "text", "Scheduled Kilimanjaro departure introduction"),
        ("hero.image", "contentImage", "Small group climbing on Kilimanjaro"),
        ("hero.primaryAction", "callToAction", "View group dates"),
        ("hero.secondaryAction", "callToAction", "Ask about a climb"),
        ("benefits", "array of object", "Small groups, Local mountain team and Scheduled departures"),
        ("benefits[].iconKey", "string", "Frontend icon identifier"),
        ("benefits[].title", "string", "Benefit title"),
        ("benefits[].description", "string", "Benefit supporting text"),
        ("schedule.eyebrow", "string", "UPCOMING DEPARTURES"),
        ("schedule.title", "string", "Best time to join a Kilimanjaro climbing group heading"),
        ("schedule.emptyTitle", "string", "No published departures message template"),
        ("schedule.emptyDescription", "text", "Private climb alternative message"),
        ("schedule.emptyAction", "callToAction", "Ask about dates"),
        ("advisor", "reference", "Victoria local trip advisor card"),
        ("reviewsTitle", "string", "Stories from the mountain"),
        ("reviewsDescription", "text", "Group-climb review introduction"),
        ("testimonials", "array of reference", "Mountain reviews displayed in the carousel"),
        ("seo", "seo", "Metadata for /groups"),
    ])
    add_body(doc, "Generated by the frontend: 2026 and 2027 year tabs; month tabs; Rain Season icon; filtered departure list; no-results month and year; join-form selected departure details.", bold_lead="Generated by the frontend:")

    add_heading(doc, "Group Departure", 1)
    add_body(doc, "Document type: groupDeparture. One document represents one scheduled climb shown in the departure list and Join Group modal.")
    add_field_table(doc, [
        ("route", "reference", "Lemosho, Machame, Rongai or Marangu route document"),
        ("durationDays", "number", "6 or 7 as displayed in the departure title"),
        ("startDate", "date", "Departure start date"),
        ("endDate", "date", "Departure end date"),
        ("highlights", "array of string", "Starry sky, Fewer hikers, Forest Birds and the other visible highlights"),
        ("price", "price", "Displayed price per person"),
        ("deposit", "price", "Displayed $100 deposit"),
        ("bookingStatus", "string list", "Scheduled, Full, Cancelled or Completed"),
        ("promotionStatus", "string list", "Standard price or Discounted price"),
        ("originalPrice", "price", "Required and visible when promotionStatus is Discounted price"),
        ("seasonStatus", "string list", "Standard season or Rain season"),
        ("seasonLabel", "string", "Required and visible when seasonStatus is Rain season"),
    ])
    add_body(doc, "Generated by the frontend: day-of-month badge, month label, date range sentence, route title, discount percentage and the values sent into the join form.", bold_lead="Generated by the frontend:")

    add_heading(doc, "FAQ Content", 1)
    add_body(doc, "The current site groups questions under Kilimanjaro, Safari, Zanzibar, Booking, Payments, Travel and General. Category headings and question content need separate records because the same questions appear on the homepage, FAQ page and trip pages.")
    add_heading(doc, "FAQ Category", 2)
    add_field_table(doc, [
        ("name", "string", "Category label such as Kilimanjaro"),
        ("sectionTitle", "string", "Climbing the mountain, Parks wildlife and itineraries, and the equivalent visible heading"),
        ("slug", "slug", "Filter identifier"),
        ("displayOrder", "number", "Order of category buttons and FAQ sections"),
    ])
    add_heading(doc, "FAQ", 2)
    add_field_table(doc, [
        ("question", "string", "Exact question shown in the accordion"),
        ("answer", "array of block", "Complete answer shown when the accordion opens"),
        ("category", "reference", "One required FAQ Category"),
    ])
    add_body(doc, "Each page stores an ordered array of FAQ references. The same FAQ can therefore appear on the FAQ page and a trip or route page without copying the answer.")

    add_heading(doc, "Testimonial", 1)
    add_body(doc, "Document type: testimonial. The seven current Tripadvisor reviews already contain the following information and every field is required.")
    add_field_table(doc, [
        ("reviewerName", "string", "Viktor, Kata, Elena P and the other visible reviewer names"),
        ("reviewDate", "date", "Structured review month and year"),
        ("travellerType", "string", "Friends, Solo or Couples"),
        ("category", "string", "Tanzania Safari, Marangu Route, Mount Kilimanjaro, Machame Route, Mount Meru or Kilimanjaro"),
        ("headline", "string", "Quoted review headline"),
        ("reviewBody", "text", "Visible review paragraph"),
        ("rating", "number", "5.0 in the existing cards"),
        ("tags", "array of string", "Three review chips displayed on each card"),
        ("reviewImage", "contentImage", "Large review image and alt text"),
        ("reviewerAvatar", "contentImage", "Tripadvisor profile image and alt text"),
        ("sourceUrl", "url", "Direct Read full review on Trip Advisor link"),
    ])

    add_heading(doc, "Local Advisor", 1)
    add_body(doc, "Document type: localAdvisor. The current Group Climbs page displays Victoria as a reusable contact card.")
    add_field_table(doc, [
        ("name", "string", "Victoria"),
        ("roleLabel", "string", "YOUR LOCAL TRIP ADVISOR"),
        ("photo", "contentImage", "Advisor portrait and alt text"),
        ("availabilityLabel", "string", "Online now"),
        ("location", "string", "Based in Tanzania"),
        ("heading", "string", "Want to plan a memorable Tanzanian adventure"),
        ("description", "text", "Local planner support message"),
        ("whatsappLabel", "string", "WhatsApp"),
        ("enquiryAction", "callToAction", "Send an enquiry"),
    ])

    add_heading(doc, "Validation Requirements", 1)
    add_field_table(doc, [
        ("All listed fields", "required validation", "Publishing is blocked when any mandatory field for the document is empty"),
        ("Slug", "custom validation", "Lowercase letters, numbers and hyphens only; unique within its document type"),
        ("Images", "nested validation", "Image asset and alternative text are both required"),
        ("Prices", "number validation", "Amount is zero or greater and currency is present"),
        ("Trip duration", "document validation", "Number of itinerary records must equal durationDays"),
        ("Route duration", "document validation", "Number of route-day records must equal itineraryDays"),
        ("Departure dates", "cross-field validation", "endDate must be later than startDate"),
        ("Discount branch", "conditional validation", "originalPrice is required and greater than price when promotionStatus is Discounted price"),
        ("Season branch", "conditional validation", "seasonLabel is required when seasonStatus is Rain season"),
        ("References", "reference validation", "Referenced accommodation, route, FAQ, testimonial and advisor documents must exist"),
        ("Arrays", "array validation", "At least one item and no duplicate references where the page visibly requires a collection"),
    ])

    add_heading(doc, "Values That Must Not Be Duplicated in Sanity", 1)
    add_bullets(doc, [
        "Day numbers are generated from itinerary array order.",
        "Trip-card category and duration filters are generated from Trip documents.",
        "Group-departure year and month filters are generated from startDate.",
        "Departure date-range text is generated from startDate and endDate.",
        "Discount percentage is calculated from originalPrice and price.",
        "Route comparison rows are generated from Kilimanjaro Route documents.",
        "Review carousel position, accommodation slider position and image counts are interface state.",
        "CSS classes, card sizes, responsive breakpoints, auto-slide timing and click behavior remain in frontend code.",
        "Sitemap last-modified dates come from Sanity document update timestamps.",
    ])

    add_heading(doc, "Studio Navigation", 1)
    add_bullets(doc, [
        "Website: Site Settings, Trips Page, Kilimanjaro Page, Kilimanjaro Shared Content and Group Climbs Page.",
        "Trips: Trips, Accommodations and Accommodation Collections.",
        "Kilimanjaro: Kilimanjaro Routes and Group Departures.",
        "Supporting Content: FAQ Categories, FAQs, Testimonials and Local Advisors.",
    ])

    add_heading(doc, "Implementation Order", 1)
    add_bullets(doc, [
        "Create shared contentImage, seo, callToAction and price object schemas.",
        "Create Site Settings and global contact data.",
        "Create Accommodation and Accommodation Collection schemas and migrate the current property library.",
        "Create Kilimanjaro Route and Kilimanjaro Shared Content schemas.",
        "Create Trip schema with itinerary days, accommodation references and inclusion groups.",
        "Create Group Departure and Group Climbs Page schemas.",
        "Create FAQ Category, FAQ, Testimonial and Local Advisor schemas.",
        "Create the Trips Page and Kilimanjaro Page singletons.",
        "Replace hardcoded data with typed GROQ queries and generate Sanity TypeGen types.",
        "Generate page metadata and the sitemap from Sanity content, then verify all existing routes before switching production traffic.",
    ])

    add_heading(doc, "Source Basis", 1)
    add_body(doc, "This specification was matched against the current application structures in src/components/reference-pages.tsx, src/components/site.tsx, src/components/home-footer.tsx, src/app/layout.tsx, src/app/[[...slug]]/page.tsx, src/app/sitemap.ts and the empty Studio schema registry in studio/schemaTypes/index.ts.")
    add_body(doc, "Schema conventions follow the installed Sanity best-practices guidance for required validation, nested objects, reusable references, ordered arrays, image hotspot and alt text, metadata generation and dynamic sitemaps.")

    configure_sections(doc)
    doc.core_properties.title = "Betheli Expeditions Sanity CMS Required Field Specification"
    doc.core_properties.subject = "Mandatory Sanity content model based on the existing Betheli Expeditions website"
    doc.core_properties.author = "Betheli Expeditions"
    doc.core_properties.keywords = "Sanity CMS, Betheli Expeditions, Next.js, content model"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
